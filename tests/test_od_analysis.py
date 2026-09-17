import csv
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from common import od_analysis as od
from common.od_map import map_html
from od_fixtures import area_fixture


class ODTests(unittest.TestCase):
    def test_gate_contract_four_classes_and_portable_od(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); source, data, zones = area_fixture(root)
            output = od.extract_trip(source, root / 'od.csv')
            records, _ = od.read_od([output])
            self.assertEqual(len(records), 4)
            # Analysis does not depend on the screening folder after extraction.
            source.rename(root / 'moved')
            result = od.analyze(records, zones, ['20260901', '20260902'])
            result['project_folder'] = '検証プロジェクト'
            self.assertEqual({k:sum(v.values()) for k,v in result['traffic'].items()}, {k:1 for k in od.TRAFFIC_TYPES})
            self.assertEqual(result['matrix'], {('A','B'):1})
            self.assertEqual(result['origins'], {'A':2}); self.assertEqual(result['destinations'], {'B':2})
            self.assertEqual(result['gate_origins'], {'G01':2}); self.assertEqual(result['gate_destinations'], {'G02':2})
            self.assertEqual(result['heat_o'], [(134.0,35.07)]*2)
            self.assertEqual(result['heat_d'], [(134.02,35.07)]*2)
            self.assertEqual(len(result['boundaries']), 2)
            out = od.export(result, root / 'out')
            book = load_workbook(out / od.result_name(result, 'OD集計.xlsx'))
            conditions = dict(book['集計条件'].values)
            self.assertEqual(conditions['ODの由来'],'トリップOD（最初行・最終行）')
            self.assertEqual(conditions['プロジェクトフォルダ名'],'検証プロジェクト')
            self.assertEqual(conditions['スクリーニングフォルダ名'],str(source.resolve()))
            for row in book['集計条件']:
                for cell in row: self.assertEqual(cell.alignment.vertical,'center')
            sheet = book['統合OD表（日平均）']
            self.assertEqual(sheet.cell(sheet.max_row,sheet.max_column).value, 2)
            self.assertEqual([sheet.cell(r,c).value for r,c in [(2,3),(2,5),(4,3),(4,5)]], [.5]*4)
            self.assertEqual(sheet.freeze_panes,'B2'); self.assertEqual(sheet['B1'].alignment.textRotation,90)
            self.assertEqual(len(sheet.conditional_formatting),1)
            self.assertEqual(book['ゲートOD明細'].max_row,4)
            self.assertEqual(book['統合OD表（全期間）']['F6'].value,4)
            self.assertFalse(list(out.glob('*.csv')))
            book.close()
            for zonal in (False, True):
                html = map_html(result, zonal)
                self.assertIn('L.circleMarker', html); self.assertIn('official_area', html)
            output.with_suffix('.context.json').unlink()
            records, _ = od.read_od([output])
            with self.assertRaisesRegex(ValueError, 'ゲートの位置情報'): od.analyze(records,zones,['20260901'])
            cells=list(od.rows(output)); cells[1][cells[0].index('o_type')]='UNKNOWN'
            od.write_csv(output,cells[0],cells[1:])
            with self.assertRaisesRegex(ValueError,'端点種別'): od.read_od([output])

    def test_gate_contract_changed_input_rejected_and_style13_keeps_real_endpoints(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp); source, data, zones=area_fixture(root)
            zips=root/'zips'; zips.mkdir()
            stream=io.StringIO(); writer=csv.writer(stream)
            for i in range(4):
                row=['']*15; row[:2]=['20260901',str(i)]; row[7]='1'; row[11:15]=['134','35.07','134.02','35.07']; writer.writerow(row)
            with zipfile.ZipFile(zips/'20260901.zip','w') as archive: archive.writestr('data.csv',stream.getvalue())
            output=od.extract(source,zips,root/'style.csv')
            records,_=od.read_od([output]); result=od.analyze(records,zones,['20260901'])
            self.assertEqual(sum(result['matrix'].values()),4)
            self.assertEqual(sum(result['gate_origins'].values()),0)
            self.assertEqual(len(result['boundaries']),2)
            self.assertEqual(od.screening_dates(source),('20260901','20260901'))
            path=data/'trip0.csv'; path.write_bytes(path.read_bytes()+b'\n')
            with self.assertRaisesRegex(ValueError,'一致しません'): od.extract_trip(data,root/'bad.csv')

    def row(self, date='20260901', opid='001', trip='1', status='OK', coords=(133.9, 35.05, 134, 35.1)):
        return dict(zip(od.FIELDS, ['fixture', date, '火', opid, trip, *coords, status, 1]))

    def test_original_trip_keys_nested_csv_and_zip(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); source = root / '15/15_area_subtrip_csv'; source.mkdir(parents=True)
            zips = root / 'zip'; zips.mkdir()
            row = [''] * 16; row[2:4] = ['20260901', '001']; row[8] = '01'
            for name in ('trip.csv', 'trip_s02.csv'):
                od.write_csv(source / name, [], [row, row])
            later = row.copy(); later[2] = '20260902'; later[8] = '1'
            od.write_csv(source / 'next.csv', [], [later])
            sr = [''] * 15; sr[0:2] = ['20260901', '001']; sr[7] = '1'; sr[11:15] = ['133.9', '35.05', '134', '35.1']
            stream = io.StringIO(); csv.writer(stream).writerow(sr)
            with zipfile.ZipFile(zips / 'unknown-name.zip', 'w') as archive:
                archive.writestr('folder/data.csv', stream.getvalue().encode('cp932'))
            output = od.extract(source.parent, zips, root / 'result.csv')
            records, duplicates = od.read_od([output, output])
            self.assertEqual(len(records), 2); self.assertEqual(duplicates, 2)
            self.assertEqual(records[0]['src_files_count'], '2')
            self.assertEqual(records[0]['status'], 'OK'); self.assertEqual(records[1]['status'], 'MISSING_OD')
            self.assertEqual(records[0]['opid'], '001')

    def test_zero_days_boundary_overlap_and_invalid_coordinates(self):
        zones = [dict(name='A', points=[(133.8,35), (133.95,35), (133.95,35.2), (133.8,35.2)]),
                 dict(name='B', points=[(133.95,35), (134.1,35), (134.1,35.2), (133.95,35.2)])]
        dates = od.target_dates('20260901', '20260903', set(range(7)))
        records = [self.row(), self.row(opid='002', coords=(133.95,35.05,135,36)),
                   self.row(opid='003', coords=(float('nan'),35,134,35)), self.row(opid='004', status='MISSING_OD')]
        result = od.analyze(records, zones, dates)
        self.assertEqual(result['days'], 3); self.assertEqual(len(result['points']), 2)
        self.assertEqual(result['matrix']['A','B'], 1)
        self.assertEqual(result['matrix']['【ゾーン重複】','【区域外】'], 1)
        self.assertEqual(result['excluded'], {'座標不正':1, 'MISSING_OD':1})
        with tempfile.TemporaryDirectory() as temp:
            output = od.export(result, Path(temp) / 'out')
            book=load_workbook(output / od.result_name(result,'OD集計.xlsx'))
            conditions=dict(book['集計条件'].values)
            self.assertEqual(conditions['ODの由来'],'様式1-3由来')
            self.assertEqual(conditions['スクリーニングフォルダ名'],'フルパス未記録（旧リスト：fixture）')
            sheet=book['統合OD表（日平均）']
            self.assertAlmostEqual(sheet.cell(sheet.max_row,sheet.max_column).value,2/3)
            book.close()
            self.assertEqual(json.loads((output / od.result_name(result, 'analysis.json')).read_text(encoding='utf-8'))['target_days'], 3)
        self.assertIn('window.updateOD', map_html(result))

    def test_conflicting_input_and_cancel(self):
        with tempfile.TemporaryDirectory() as temp:
            file = Path(temp) / 'od.csv'
            od.write_csv(file, od.FIELDS, [self.row().values(), self.row(coords=(130,35,134,35)).values()])
            with self.assertRaisesRegex(ValueError, '異なるOD'): od.read_od([file])
            with self.assertRaises(InterruptedError): od.read_od([file], cancel=lambda: True)

    def test_zone_csv_and_weekdays(self):
        with tempfile.TemporaryDirectory() as temp:
            file = Path(temp) / 'zones.csv'
            od.write_csv(file, ['A',133,35,134,35,134,36], [])
            self.assertEqual(len(od.load_zones(temp)),1)
        self.assertEqual(od.target_dates('20260901','20260907',{0}),['20260907'])
        with self.assertRaises(ValueError): od.target_dates('20260907','20260901',{0})

    def test_trip_endpoints_multiple_trips_clipped_segments_and_duplicates(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); source = root/'input'; source.mkdir()
            def row(trip, lon, lat):
                r=['']*16; r[2:4]=['20260901','00001']; r[8]=trip; r[14:16]=[str(lon),str(lat)]; return r
            part1=[row('1',133.9,35.01),row('1',133.95,35.02),row('2',134,35.03)]
            part2=[row('1',134.1,35.1),row('1',134.2,35.2)]
            od.write_csv(source/'first.csv',[],part1); od.write_csv(source/'copy.csv',[],part1)
            od.write_csv(source/'subtrip_s02.csv',[],part2)
            output=od.extract_trip(source,root/'trip.csv')
            records,duplicates=od.read_od([output,output])
            self.assertEqual(len(records),3); self.assertEqual(duplicates,3)
            endpoints={od.coordinates(r) for r in records}
            self.assertEqual(endpoints,{(133.9,35.01,133.95,35.02),(134,35.03,134,35.03),(134.1,35.1,134.2,35.2)})
            self.assertEqual(od.method_of(records),'trip')
            self.assertEqual(sum(r['src_files_count']=='2' for r in records),2)
            self.assertEqual(len({r['trip_instance'] for r in records}),3)

    def test_trip_invalid_first_row_is_not_replaced_and_modes_cannot_mix(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp); source=root/'input'; source.mkdir()
            r=['']*16; r[2:4]=['20260901','001']; r[8]='1'; r[14:16]=['','35']
            end=r.copy(); end[14]='134'
            od.write_csv(source/'a.csv',[],[r,end])
            trip=od.extract_trip(source,root/'trip.csv'); records,_=od.read_od([trip])
            self.assertEqual(records[0]['status'],'INVALID_COORDINATES'); self.assertEqual(records[0]['o_lon'],'')
            style=root/'style.csv'; od.write_csv(style,od.FIELDS,[self.row().values()])
            with self.assertRaisesRegex(ValueError,'異なるOD方式'): od.read_od([trip,style])


if __name__ == '__main__': unittest.main()
