"""Japanese Excel workbook for one route-path analysis result."""
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CLASS_NAMES = {
    'ALL': '全交通', 'THROUGH': '通過交通',
    'EXTERNAL_TO_INTERNAL': '外内交通', 'INTERNAL_TO_EXTERNAL': '内外交通',
    'INTERNAL': '内々交通', 'UNKNOWN': '分類不明',
}


def write_route_workbook(path, *, project, target, counts, labels, matrix, ranking,
                         classified, mesh_rows, official, warnings):
    """Write the CSV-equivalent route results into one readable Japanese workbook."""
    book = Workbook()
    conditions = book.active
    conditions.title = '集計条件'
    conditions.append(['項目', '内容'])
    conditions_rows = [
        ('プロジェクトフォルダ', str(project)),
        ('対象路線', target.name),
        ('第2スクリーニングデータ', str(target.folder)),
        ('対象トリップ数', counts['ALL']),
        ('正式値', '有効' if official else '無効（警告を確認）'),
        ('OD行', 'O：起点・出発'),
        ('OD列', 'D：終点・到着'),
        ('内側端点', '12_ゾーニングデータのポリゴンで名称を判定'),
        ('境界端点', '14_エリアデータの指定ゲートへ割当'),
    ]
    conditions_rows.extend(('警告', warning) for warning in sorted(set(warnings)))
    for row in conditions_rows:
        conditions.append(row)

    matrix_sheet = book.create_sheet('統合OD表（全期間）')
    matrix_sheet.append(['O \\ D', *labels, '合計'])
    for origin in labels:
        matrix_sheet.append([origin, *(matrix[origin, destination] for destination in labels),
                             sum(matrix[origin, destination] for destination in labels)])
    matrix_sheet.append(['合計', *(sum(matrix[origin, destination] for origin in labels) for destination in labels),
                         sum(matrix.values())])
    matrix_sheet.freeze_panes = 'B2'
    matrix_sheet.row_dimensions[1].height = 150
    matrix_sheet.column_dimensions['A'].width = 30
    for column in range(2, len(labels) + 3):
        matrix_sheet.column_dimensions[get_column_letter(column)].width = 11
        matrix_sheet.cell(1, column).alignment = Alignment(text_rotation=90, horizontal='center', vertical='center')
    if labels:
        matrix_sheet.conditional_formatting.add(
            f'B2:{get_column_letter(len(labels) + 1)}{len(labels) + 1}',
            ColorScaleRule(start_type='num', start_value=0, start_color='F2F8FC',
                           end_type='max', end_color='20769B'))

    summary = book.create_sheet('交通区分別集計')
    summary.append(['交通区分', 'トリップ数', '全交通に占める割合'])
    total = counts['ALL']
    for key in ('ALL', 'THROUGH', 'EXTERNAL_TO_INTERNAL', 'INTERNAL_TO_EXTERNAL', 'INTERNAL', 'UNKNOWN'):
        summary.append([CLASS_NAMES[key], counts[key], counts[key] / total if total else 0])

    gate_od = book.create_sheet('ゲートOD明細')
    gate_od.append(['起点ゲート', '終点ゲート', 'トリップ数', '通過交通内割合', '全交通内割合'])
    for row in ranking:
        gate_od.append([row['start_gate'], row['end_gate'], row['trip_count'],
                        row['share_within_class'], row['share_total']])

    trips = book.create_sheet('トリップ分類一覧')
    trips.append(['トリップID', '元CSV', '対象路線', '起点種別', '起点ゲート', '終点種別', '終点ゲート',
                  '起点経度', '起点緯度', '終点経度', '終点緯度', '起点OD', '終点OD', '交通区分', '分類グループ'])
    for row in classified:
        trips.append([row.get(key, '') for key in ('trip_id', 'source_file', 'route_name', 'start_type', 'start_gate_id',
                     'end_type', 'end_gate_id', 'start_lon', 'start_lat', 'end_lon', 'end_lat', 'origin',
                     'destination', 'od_class', 'destination_group')])

    meshes = book.create_sheet('メッシュ集計')
    meshes.append(['表示グループ', 'メッシュX', 'メッシュY', 'トリップ数', '表示群内割合', '原点経度', '原点緯度', 'メッシュサイズ（m）'])
    for row in mesh_rows:
        meshes.append([row[key] for key in ('group', 'cell_x', 'cell_y', 'trip_count', 'share',
                       'origin_lon', 'origin_lat', 'cell_size_m')])

    for sheet in book:
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='23465D')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(vertical='center')
        if sheet is matrix_sheet:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.row == len(labels) + 2 or cell.column == len(labels) + 2:
                        cell.fill = PatternFill('solid', fgColor='D7E6F0')
                        cell.font = Font(bold=True, color='163047')
            continue
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 18
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        if sheet is conditions:
            sheet.column_dimensions['A'].width = 28
            sheet.column_dimensions['B'].width = 95
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='center')
                sheet.row_dimensions[row[0].row].height = 30 * (str(row[1].value).count('\n') + 1)
        if sheet in (summary, gate_od, meshes):
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = '0.0%' if cell.column in (3, 4, 5) else '#,##0.000000'
    book.save(path)
