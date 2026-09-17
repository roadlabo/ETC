"""One Japanese workbook per OD analysis, including the complete zone/gate matrix."""
from pathlib import Path, PureWindowsPath
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter


def write_workbook(result, path):
    from common.od_analysis import combined_matrix, display_label, METHODS
    workbook = Workbook()
    summary = workbook.active; summary.title = '集計条件'
    for row in [('項目', '内容'), ('OD方式', METHODS[result['method']]),
                ('ODの由来', '様式1-3由来' if result['method'] == 'style13' else 'トリップOD（最初行・最終行）'),
                ('プロジェクトフォルダ名', result.get('project_folder') or '未記録'),
                ('スクリーニングフォルダ名', '\n'.join(
                    folder if Path(folder).is_absolute() or PureWindowsPath(folder).is_absolute()
                    else 'フルパス未記録（旧リスト：' + folder + '）'
                    for folder in result.get('screening_folders', [])) or '未記録'),
                ('対象日数', result['days']), ('採用トリップ数', len(result['points'])),
                ('行方向', 'O：起点・発生'), ('列方向', 'D：終点・集中'),
                ('統合OD表', 'ゾーン→ゾーン・ゾーン→ゲート・ゲート→ゾーン・ゲート→ゲートをすべて含む'),
                ('地図の分布', '非ゲート端点のみ。ゲート端点は別の円で表示。')]: summary.append(row)
    for kind, count in result['excluded'].items(): summary.append(('除外：' + kind, count))
    labels, matrix = combined_matrix(result)
    n = len(labels)
    for title, divisor in [('統合OD表（全期間）', 1), ('統合OD表（日平均）', result['days'])]:
        sheet = workbook.create_sheet(title)
        sheet.append(['行＝O（起点）／列＝D（終点）', *[display_label(s) for s in labels], '合計'])
        for a in labels:
            sheet.append([display_label(a), *[matrix[a, b] / divisor for b in labels], sum(matrix[a, b] for b in labels) / divisor])
        sheet.append(['合計', *[sum(matrix[a, b] for a in labels) / divisor for b in labels], sum(matrix.values()) / divisor])
        sheet.freeze_panes = 'B2'; sheet.row_dimensions[1].height = 150
        sheet.column_dimensions['A'].width = 30
        for col in range(2, n + 3):
            sheet.column_dimensions[get_column_letter(col)].width = 10
            sheet.cell(1, col).alignment = Alignment(text_rotation=90, horizontal='center', vertical='center')
        for row in sheet.iter_rows(min_row=2, min_col=2):
            for cell in row: cell.number_format = '#,##0' if divisor == 1 and title.endswith('全期間）') else '#,##0.00'
        if n:
            sheet.conditional_formatting.add(f'B2:{get_column_letter(n+1)}{n+1}', ColorScaleRule(
                start_type='num', start_value=0, start_color='F2F8FC', end_type='max', end_color='20769B'))
        # Keep totals visually separate from the cell colour scale.
        for row in sheet.iter_rows():
            for cell in row:
                if cell.row == n + 2 or cell.column == n + 2:
                    cell.fill = PatternFill('solid', fgColor='D7E6F0'); cell.font = Font(bold=True, color='163047')
        if result['labels'] and result['gates']:
            boundary = len(result['labels']) + 2
            edge = Side(style='medium', color='C58632')
            for row in range(1, n + 3): sheet.cell(row, boundary).border = Border(left=edge)
            for col in range(1, n + 3):
                cell = sheet.cell(boundary, col)
                cell.border = Border(left=cell.border.left, top=edge)
        sheet.print_title_rows = '1:1'; sheet.print_title_cols = 'A:A'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    sheet = workbook.create_sheet('交通区分別集計'); sheet.append(['交通区分','全期間（トリップ）','日平均（トリップ/日）'])
    for kind, counts in result['traffic'].items(): sheet.append([kind,sum(counts.values()),sum(counts.values())/result['days']])
    for title, entities, origins, destinations in [
        ('ゾーン別発生集中', [(s,s) for s in result['labels']], result['origins'], result['destinations']),
        ('ゲート別発生集中', [(g['gate_id'], g['name']) for g in result['gates']], result['gate_origins'], result['gate_destinations'])]:
        sheet = workbook.create_sheet(title)
        sheet.append(['番号・ゾーン','名称','発生（全期間）','集中（全期間）','発生（日平均）','集中（日平均）'])
        for identity, name in entities:
            sheet.append([identity,name,origins[identity],destinations[identity],origins[identity]/result['days'],destinations[identity]/result['days']])
    sheet = workbook.create_sheet('ゲートOD明細'); sheet.append(['交通区分','O：起点','D：終点','全期間（トリップ）','日平均（トリップ/日）'])
    for kind, counts in result['traffic'].items():
        if kind != '内々交通':
            for (a,b), count in sorted(counts.items()): sheet.append([kind,display_label(a),display_label(b),count,count/result['days']])
    sheet = workbook.create_sheet('対象日'); sheet.append(['対象運行日'])
    for date in result['dates']: sheet.append([date])
    for sheet in workbook:
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='23465D'); cell.font = Font(color='FFFFFF', bold=True)
        if not sheet.title.startswith('統合OD表'):
            sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
            for col in range(1, sheet.max_column + 1): sheet.column_dimensions[get_column_letter(col)].width = 25
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    # Treat names as text even when users begin them with '='.
                    if isinstance(cell.value, str): cell.data_type = 's'
                    elif isinstance(cell.value, float): cell.number_format = '#,##0.00'
            if sheet.title == '集計条件':
                sheet.column_dimensions['B'].width = 95
                for row in sheet.iter_rows(min_row=2):
                    for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='center')
                    sheet.row_dimensions[row[0].row].height = 30 * (str(row[1].value).count('\n') + 1)
                for cell in sheet[1]: cell.alignment = Alignment(vertical='center')
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
    workbook.save(path)
